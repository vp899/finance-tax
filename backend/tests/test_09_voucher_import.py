"""凭证导入（新格式：凭证类别 凭证号 凭证日期 附单据数 摘要 科目编码 …）测试（跑在独立账套 t09）"""
import io

import pytest
from openpyxl import Workbook, load_workbook

from conftest import api, delete, get, post, put, upload

B = "book_id=t09"


def p(path):
    return path + ("&" if "?" in path else "?") + B


HEADERS = [
    "凭证类别", "凭证号", "凭证日期", "附单据数", "摘要", "科目编码", "科目名称",
    "借方金额", "贷方金额",
    "项目编码", "项目", "客户编码", "客户", "供应商编码", "供应商",
    "部门编码", "部门", "员工编码", "员工", "存货编码", "存货",
    "规格型号", "数量", "计量单位", "单价", "外币金额", "币种", "汇率", "制单人", "审核人",
]


def row(no, date, summary, code, name, debit, credit, **kw):
    """按新模板列生成一行（可用 kw 覆盖任意列名）"""
    values = {
        "凭证类别": "记", "凭证号": no, "凭证日期": date, "附单据数": 1,
        "摘要": summary, "科目编码": code, "科目名称": name,
        "借方金额": debit, "贷方金额": credit,
        "币种": "CNY", "汇率": 1, "制单人": "张三", "审核人": "李四",
    }
    values.update(kw)
    return [values.get(h, "") for h in HEADERS]


def xlsx(rows, headers=HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.title = "凭证导入"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module", autouse=True)
def book(server):
    delete("/api/booksets/t09")
    st, r = post("/api/booksets", {"id": "t09", "name": "凭证导入测试账套", "opening_year": "2033"})
    assert st == 200
    yield
    delete("/api/booksets/t09")


class TestTemplate:
    def test_模板为新格式30列(self):
        st, data = api("GET", "/api/data/template/vouchers", raw=True)
        assert st == 200
        ws = load_workbook(io.BytesIO(data)).active
        headers = [c.value for c in ws[1]]
        assert headers == HEADERS


class TestNewFormat:
    def test_导入并保留凭证号与人员(self):
        data = xlsx([
            row("记-101", "2033-01-05", "提现备用金", "1001", "库存现金", 5000, 0),
            row("记-101", "2033-01-05", "提现备用金", "1002", "银行存款", 0, 5000),
        ])
        st, r = upload(p("/api/data/import/vouchers"), "v.xlsx", data)
        assert st == 200, r
        assert r["created"] == 1 and r["errors"] == []
        rows = get(p("/api/vouchers?period=2033-01"))[1]["rows"]
        v = rows[0]
        assert v["source_no"] == "记-101"
        assert v["maker"] == "张三" and v["reviewer"] == "李四"
        assert v["attachment_count"] == 1
        assert v["status"] == "posted"

    def test_明细字段落库(self):
        data = xlsx([
            row("记-102", "2033-01-31", "采购办公用品", "5602", "管理费用", 300, 0,
                项目编码="P001", 项目="装修项目", 部门编码="D01", 部门="办公室",
                规格型号="A4 纸", 数量=10, 计量单位="箱", 单价=30,
                外币金额="", 币种="CNY", 汇率=1),
            row("记-102", "2033-01-31", "采购办公用品", "1001", "库存现金", 0, 300),
        ])
        st, r = upload(p("/api/data/import/vouchers"), "v2.xlsx", data)
        assert st == 200, r
        vid = get(p("/api/vouchers?period=2033-01&size=50"))[1]["rows"][0]["id"]
        detail = get(p(f"/api/vouchers/{vid}"))[1]
        e = next(x for x in detail["entries"] if x["account_code"] == "5602")
        assert e["spec"] == "A4 纸" and e["quantity"] == 10 and e["unit"] == "箱"
        assert e["price"] == 30
        assert '"project"' in e["aux"] and "装修项目" in e["aux"]
        assert '"dept"' in e["aux"] and "办公室" in e["aux"]

    def test_凭证号向下补全(self):
        data = xlsx([
            row("记-103", "2033-02-10", "费用报销", "5602", "管理费用", 100, 0),
            row("", "2033-02-10", "费用报销", "1001", "库存现金", 0, 100),  # 续行留空
        ])
        st, r = upload(p("/api/data/import/vouchers"), "v3.xlsx", data)
        assert st == 200, r
        assert r["created"] == 1
        v = get(p("/api/vouchers?period=2033-02"))[1]["rows"][0]
        assert v["source_no"] == "记-103" and v["entry_count"] == 2

    def test_重复导入幂等(self):
        data = xlsx([
            row("记-104", "2033-03-05", "重复导入", "1001", "库存现金", 200, 0),
            row("记-104", "2033-03-05", "重复导入", "1002", "银行存款", 0, 200),
        ])
        st, r1 = upload(p("/api/data/import/vouchers"), "v4.xlsx", data)
        assert st == 200 and r1["created"] == 1
        st, r2 = upload(p("/api/data/import/vouchers"), "v4.xlsx", data)
        assert st == 200 and r2["created"] == 0 and r2["skipped"] == 1
        assert get(p("/api/vouchers?period=2033-03"))[1]["total"] == 1

    def test_补齐草稿凭证对方科目后转正式(self):
        """明细账单侧导入生成的草稿，可用凭证导入补另一侧自动转正式"""
        detail_rows = [
            [1, "5602", "管理费用", "2033-04-05", "记-105", "分批导入", 60.00, None, None, None],
        ]
        wb = Workbook(); ws = wb.active; ws.title = "明细账"
        ws.append(["序号", "科目编码", "科目", "日期", "凭证号", "摘要", "借方", "贷方", "方向", "余额"])
        for r in detail_rows:
            ws.append(r)
        buf = io.BytesIO(); wb.save(buf)
        st, r = upload(p("/api/data/import/detail-ledger"), "d1.xlsx", buf.getvalue())
        assert st == 200 and r["draft_vouchers"] == 1
        # 凭证导入补对方科目
        data = xlsx([row("记-105", "2033-04-05", "分批导入", "1001", "库存现金", 0, 60)])
        st, r = upload(p("/api/data/import/vouchers"), "p2.xlsx", data)
        assert st == 200 and r["merged"] == 1, r
        v = next(x for x in get(p("/api/vouchers?period=2033-04"))[1]["rows"]
                 if x["source_no"] == "记-105")
        assert v["entry_count"] == 2 and v["status"] == "posted"

    def test_合并后不平被拒(self):
        # 同凭证号再导一行单边分录 → 合并后不平 → 拒绝并保留原凭证
        data = xlsx([row("记-105", "2033-04-05", "多余行", "5602", "管理费用", 999, 0)])
        st, r = upload(p("/api/data/import/vouchers"), "p3.xlsx", data)
        assert st == 400 and "合并后借贷不平衡" in r["detail"]
        v = next(x for x in get(p("/api/vouchers?period=2033-04"))[1]["rows"]
                 if x["source_no"] == "记-105")
        assert v["entry_count"] == 2 and v["total_debit"] == 60

    def test_导出凭证明细含新列(self):
        st, data = api("GET", p("/api/data/export/vouchers?from_period=2033-01&to_period=2033-04"), raw=True)
        assert st == 200
        ws = load_workbook(io.BytesIO(data)).active
        headers = [c.value for c in ws[1]]
        for col in ("凭证号", "借方金额", "贷方金额", "规格型号", "单价", "外币金额",
                    "项目编码", "项目", "部门编码", "部门", "附单据数", "制单人", "审核人"):
            assert col in headers
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        idx = {h: i for i, h in enumerate(headers)}
        r = next(x for x in rows if str(x[idx["摘要"]]) == "采购办公用品"
                 and str(x[idx["科目编码"]]) == "5602")
        assert r[idx["规格型号"]] == "A4 纸"
        assert r[idx["项目"]] == "装修项目"
        assert r[idx["制单人"]] == "张三"

    def test_导出文件可直接再导入(self):
        st, data = api("GET", p("/api/data/export/vouchers?from_period=2033-01&to_period=2033-04"), raw=True)
        st, r = upload(p("/api/data/import/vouchers"), "roundtrip.xlsx", data)
        assert st == 200
        assert r["created"] == 0          # 凭证号匹配 → 全部识别为已存在
        assert r["skipped"] >= 1


class TestLegacyAndErrors:
    def test_旧格式八列模板仍可导入(self):
        old_headers = ["日期(YYYY-MM-DD)", "凭证类型", "摘要", "科目编码", "借方金额", "贷方金额", "数量", "单位"]
        rows = [["2033-05-05", "记", "旧格式-现金", "1001", 800, 0, "", ""],
                ["2033-05-05", "记", "旧格式-现金", "1002", 0, 800, "", ""]]
        st, r = upload(p("/api/data/import/vouchers"), "old.xlsx", xlsx(rows, old_headers))
        assert st == 200 and r["created"] == 1
        v = get(p("/api/vouchers?period=2033-05"))[1]["rows"][0]
        assert v["total_debit"] == 800 and v["voucher_no"].startswith("记-")

    def test_缺日期或科目报错(self):
        data = xlsx([row("记-201", "", "缺日期", "1001", "库存现金", 10, 0)])
        st, r = upload(p("/api/data/import/vouchers"), "bad1.xlsx", data)
        assert st == 400 and "日期或科目编码为空" in r["detail"]

    def test_科目不存在报错(self):
        data = xlsx([
            row("记-202", "2033-06-05", "错误科目", "999999", "不存在", 10, 0),
            row("记-202", "2033-06-05", "错误科目", "1001", "库存现金", 0, 10),
        ])
        st, r = upload(p("/api/data/import/vouchers"), "bad2.xlsx", data)
        assert st == 400 and "不存在" in r["detail"]

    def test_借贷不平衡报错(self):
        data = xlsx([row("记-203", "2033-06-06", "不平衡", "1001", "库存现金", 10, 0)])
        st, r = upload(p("/api/data/import/vouchers"), "bad3.xlsx", data)
        assert st == 400 and "不平衡" in r["detail"]

    def test_新格式缺凭证号报错(self):
        data = xlsx([
            row("", "2033-06-07", "缺凭证号", "1001", "库存现金", 10, 0),
            row("", "2033-06-07", "缺凭证号", "1002", "银行存款", 0, 10),
        ])
        st, r = upload(p("/api/data/import/vouchers"), "bad4.xlsx", data)
        assert st == 400 and "凭证号为空" in r["detail"]

    def test_负数金额报错(self):
        data = xlsx([
            row("记-204", "2033-06-08", "负数", "1001", "库存现金", -10, 0),
            row("记-204", "2033-06-08", "负数", "1002", "银行存款", 0, -10),
        ])
        st, r = upload(p("/api/data/import/vouchers"), "bad5.xlsx", data)
        assert st == 400
