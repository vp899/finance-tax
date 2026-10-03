"""导入其它平台凭证（含结转类凭证）后的余额与报表口径测试（跑在独立账套 t11）

场景（与手工推演一致）：
  期初：银行存款 200000（借）/ 实收资本 200000（贷），会计年度 2034
  2034-01：购货 20000；销售 51500（收入 50000 + 税 1500）；结转销售成本 12000；
          结转本期损益（收入 50000 / 成本 12000 → 本年利润）
  2034-02：销售 206000（收入 200000 + 税 6000）；购货 60000；结转销售成本 40000；
          结转本期损益（收入 200000 / 成本 40000 → 本年利润）
  2034-12：结转本年利润至未分配利润 198000

预期（修复后）：
  - 期初/期末余额按余额口径列示（单边净额），期初+本期=期末
  - 结转类凭证被识别并登记结转记录（利润表口径剔除结转损益凭证）
  - 利润表/资产负债表金额与手工推演一致，试算平衡表三组合计均平衡
"""
import io

import pytest
from openpyxl import Workbook, load_workbook

from conftest import api, delete, get, post, put, upload

B = "book_id=t11"


def p(path):
    return path + ("&" if "?" in path else "?") + B


HEADERS = ["凭证类别", "凭证号", "凭证日期", "附单据数", "摘要", "科目编码", "科目名称",
           "借方金额", "贷方金额",
           "项目编码", "项目", "客户编码", "客户", "供应商编码", "供应商",
           "部门编码", "部门", "员工编码", "员工", "存货编码", "存货",
           "规格型号", "数量", "计量单位", "单价", "外币金额", "币种", "汇率", "制单人", "审核人"]


def row(no, date, summary, code, name, debit, credit, **kw):
    values = {"凭证类别": "记", "凭证号": no, "凭证日期": date, "附单据数": 1,
              "摘要": summary, "科目编码": code, "科目名称": name,
              "借方金额": debit, "贷方金额": credit,
              "币种": "CNY", "汇率": 1, "制单人": "张三", "审核人": "李四"}
    values.update(kw)
    return [values.get(h, "") for h in HEADERS]


def xlsx(rows, headers=HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.title = "凭证导入"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DETAIL_HEADERS = ["序号", "科目编码", "科目", "日期", "凭证号", "摘要",
                  "借方", "贷方", "方向", "余额"]


IMPORT_ROWS = [
    row("记-001", "2034-01-05", "购进商品", "1405", "库存商品", 20000, 0),
    row("记-001", "2034-01-05", "购进商品", "2202", "应付账款", 0, 20000),
    row("记-002", "2034-01-10", "销售商品", "1002", "银行存款", 51500, 0),
    row("记-002", "2034-01-10", "销售商品", "5001", "主营业务收入", 0, 50000),
    row("记-002", "2034-01-10", "销售商品", "222101", "应交增值税", 0, 1500),
    row("记-021", "2034-01-31", "结转本月销售成本", "5401", "主营业务成本", 12000, 0),
    row("记-021", "2034-01-31", "结转本月销售成本", "1405", "库存商品", 0, 12000),
    row("记-022", "2034-01-31", "结转本期损益", "5001", "主营业务收入", 50000, 0),
    row("记-022", "2034-01-31", "结转本期损益", "3103", "本年利润", 0, 50000),
    row("记-022", "2034-01-31", "结转本期损益", "3103", "本年利润", 12000, 0),
    row("记-022", "2034-01-31", "结转本期损益", "5401", "主营业务成本", 0, 12000),
    row("记-001", "2034-02-10", "销售商品", "1002", "银行存款", 206000, 0),
    row("记-001", "2034-02-10", "销售商品", "5001", "主营业务收入", 0, 200000),
    row("记-001", "2034-02-10", "销售商品", "222101", "应交增值税", 0, 6000),
    row("记-002", "2034-02-20", "购进商品", "1405", "库存商品", 60000, 0),
    row("记-002", "2034-02-20", "购进商品", "2202", "应付账款", 0, 60000),
    row("记-021", "2034-02-28", "结转本月销售成本", "5401", "主营业务成本", 40000, 0),
    row("记-021", "2034-02-28", "结转本月销售成本", "1405", "库存商品", 0, 40000),
    row("记-022", "2034-02-28", "结转本期损益", "5001", "主营业务收入", 200000, 0),
    row("记-022", "2034-02-28", "结转本期损益", "3103", "本年利润", 0, 200000),
    row("记-022", "2034-02-28", "结转本期损益", "3103", "本年利润", 40000, 0),
    row("记-022", "2034-02-28", "结转本期损益", "5401", "主营业务成本", 0, 40000),
    row("记-121", "2034-12-31", "结转本年利润至未分配利润", "3103", "本年利润", 198000, 0),
    row("记-121", "2034-12-31", "结转本年利润至未分配利润", "3104", "利润分配-未分配利润", 0, 198000),
]


def acc_id(code):
    for a in get(p("/api/accounts"))[1]:
        if a["code"] == code:
            return a["id"]
    raise AssertionError(code)


@pytest.fixture(scope="module", autouse=True)
def book(server):
    delete("/api/booksets/t11")
    st, r = post("/api/booksets", {"id": "t11", "name": "导入结转测试账套", "opening_year": "2034"})
    assert st == 200
    assert put(p("/api/accounts/openings/save"), {
        "year": "2034",
        "rows": [{"account_id": acc_id("1002"), "debit": 200000},
                 {"account_id": acc_id("3001"), "credit": 200000}]})[0] == 200
    st, r = upload(p("/api/data/import/vouchers"), "other_platform.xlsx", xlsx(IMPORT_ROWS))
    assert st == 200, r
    assert r["errors"] == []
    yield
    delete("/api/booksets/t11")


def _tb(period=None, **kw):
    q = f"period={period}" if period else "&".join(f"{k}={v}" for k, v in kw.items())
    return get(p(f"/api/books/trial-balance?{q}"))[1]


def _rows_by_code(data):
    return {x["code"]: x for x in data["rows"]}


class TestCarryoverRecognized:
    def test_导入识别结转类凭证(self):
        st, r = upload(p("/api/data/import/vouchers"), "again.xlsx", xlsx(IMPORT_ROWS))
        assert st == 200, r
        assert r["created"] == 0 and r["merged"] == 0 and r["skipped"] == 9
        assert r["carryover_marked"] == 0 and r["carryover_records"] == 0  # 幂等，不重复登记

    def test_结转记录按类型登记(self):
        recs = get(p("/api/carryover/records?year=2034"))[1]
        kinds = sorted(x["kind"] for x in recs)
        assert kinds == ["profit", "profit", "retain_profit", "sales_cost", "sales_cost"]
        assert all(x["status"] == "active" for x in recs)
        profit = [x for x in recs if x["kind"] == "profit"]
        assert sorted(x["amount"] for x in profit) == [62000, 240000]

    def test_凭证携带结转类型(self):
        rows = get(p("/api/vouchers?from_period=2034-01&to_period=2034-12&size=50"))[1]["rows"]
        kinds = {x["source_no"]: x for x in rows}
        v = next(x for x in rows if x["source_no"] == "记-022" and x["period"] == "2034-01")
        detail = get(p(f"/api/vouchers/{v['id']}"))[1]
        assert detail.get("carryover_kind") == "profit"

    def test_业务凭证不被误判为结转(self):
        rows = get(p("/api/vouchers?from_period=2034-01&to_period=2034-12&size=50"))[1]["rows"]
        for x in rows:
            if x["source_no"] in ("记-001", "记-002"):
                assert get(p(f"/api/vouchers/{x['id']}"))[1].get("carryover_kind") in (None, "")


class TestOpeningBalance:
    def test_试算平衡期初为余额口径(self):
        rows = _rows_by_code(_tb("2034-02"))
        # 结转后损益类/本年利润期初应是余额，而不是区间前的借贷累计发生额
        assert rows["5001"]["opening_debit"] == 0 and rows["5001"]["opening_credit"] == 0
        assert rows["5401"]["opening_debit"] == 0 and rows["5401"]["opening_credit"] == 0
        assert rows["3103"]["opening_debit"] == 0 and rows["3103"]["opening_credit"] == 38000
        assert rows["1405"]["opening_debit"] == 8000 and rows["1405"]["opening_credit"] == 0
        assert rows["1002"]["opening_debit"] == 251500 and rows["1002"]["opening_credit"] == 0

    def test_期初加本期等于期末(self):
        for period in ("2034-01", "2034-02", "2034-12"):
            for x in _tb(period)["rows"]:
                od, oc = x["opening_debit"], x["opening_credit"]
                pd, pc = x["period_debit"], x["period_credit"]
                cd, cc = x["closing_debit"], x["closing_credit"]
                net = round((od + pd) - (oc + pc), 2)
                assert (cd, cc) == ((net, 0.0) if net >= 0 else (0.0, -net)), (period, x["code"])

    def test_试算三组合计均平衡(self):
        tb = _tb("2034-02")
        assert tb["total_opening_debit"] == tb["total_opening_credit"] == 259500
        assert tb["total_period_debit"] == tb["total_period_credit"] == 546000
        assert tb["total_debit"] == tb["total_credit"] == 485500
        assert tb["balanced"] is True

    def test_总账与余额表期初口径一致(self):
        gl = _rows_by_code(get(p("/api/books/general-ledger?period=2034-02"))[1])
        bt = _rows_by_code(get(p("/api/books/balance-table?period=2034-02"))[1])
        for code in ("1002", "1405", "5001", "5401", "3103"):
            assert gl[code]["opening_debit"] == bt[code]["opening_debit"]
            assert gl[code]["opening_credit"] == bt[code]["opening_credit"]
        assert gl["3103"]["opening_credit"] == 38000

    def test_明细账期初余额口径(self):
        d = get(p("/api/books/detail?account_code=5001&from_period=2034-02&to_period=2034-02"))[1]
        assert d["opening"] == {"debit": 0, "credit": 0}
        d2 = get(p("/api/books/detail?account_code=3103&from_period=2034-02&to_period=2034-02"))[1]
        assert d2["opening"] == {"debit": 0, "credit": 38000}


class TestReports:
    def test_利润表剔除结转损益后金额正确(self):
        r = get(p("/api/reports/income?period=2034-02&mode=month"))[1]
        m = {x["name"]: x for x in r["rows"]}
        assert m["一、营业收入"]["current"] == 200000
        assert m["一、营业收入"]["ytd"] == 250000
        assert m["　减：营业成本"]["current"] == 40000
        assert m["　减：营业成本"]["ytd"] == 52000
        assert m["四、净利润（净亏损以“-”号填列）"]["current"] == 160000
        assert m["四、净利润（净亏损以“-”号填列）"]["ytd"] == 198000

    def test_利润表月报一月金额(self):
        r = get(p("/api/reports/income?period=2034-01&mode=month"))[1]
        m = {x["name"]: x for x in r["rows"]}
        assert m["一、营业收入"]["current"] == 50000
        assert m["四、净利润（净亏损以“-”号填列）"]["current"] == 38000

    def test_资产负债表勾稽与未分配利润(self):
        r = get(p("/api/reports/balance-sheet?period=2034-02"))[1]
        m = {x["name"]: x for x in r["rows"]}
        assert r["balanced"] is True
        assert m["未分配利润"]["ending"] == 198000
        assert m["未分配利润"]["beginning"] == 0
        assert m["资产总计"]["ending"] == 485500
        assert abs(m["资产总计"]["ending"] - m["负债和所有者权益总计"]["ending"]) < 0.01

    def test_资产负债表年末未分配利润结转后仍正确(self):
        r = get(p("/api/reports/balance-sheet?period=2034-12"))[1]
        m = {x["name"]: x for x in r["rows"]}
        assert r["balanced"] is True
        assert m["未分配利润"]["ending"] == 198000  # 已结转到 3104

    def test_现金流量表期末现金等于账面(self):
        r = get(p("/api/reports/cashflow?from_period=2034-01&to_period=2034-02"))[1]
        assert r["balanced"] is True


class TestClose:
    def test_导入数据可正常结账(self):
        st, r = post(p("/api/carryover/close"), {"period": "2034-01"})
        assert st == 200 and r["status"] == "closed"

    def test_结账后可反结账(self):
        assert post(p("/api/carryover/open"), {"period": "2034-01"})[0] == 200


class TestDetailLedgerRoundTrip:
    def test_明细账导出期初按余额单边列示(self):
        st, data = api("GET", p("/api/data/export/detail-ledger"
                                "?from_period=2034-02&to_period=2034-02"), raw=True)
        assert st == 200
        ws = load_workbook(io.BytesIO(data)).active
        rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
        opening = [r for r in rows if str(r[4]) == "期初余额"]
        # 期初行不允许借贷同时有值（余额口径单边）
        for r in opening:
            assert not (r[6] and r[7]), r
        r5001 = next(r for r in opening if str(r[1]) == "5001")
        assert (r5001[6], r5001[7]) == (None, None)  # 结转后期初为平

    def test_导出导入往返无错误(self):
        st, data = api("GET", p("/api/data/export/detail-ledger"
                                "?from_period=2034-01&to_period=2034-12"), raw=True)
        st, r = upload(p("/api/data/import/detail-ledger"), "roundtrip.xlsx", data)
        assert st == 200, r
        assert r["errors"] == [] and r["warnings"] == []
        assert r["created_vouchers"] == 0 and r["skipped_vouchers"] == 9

    def test_期初借贷同时有值按净额处理(self):
        delete("/api/booksets/t11x")
        assert post("/api/booksets", {"id": "t11x", "name": "期初净额", "opening_year": "2035"})[0] == 200
        rows = [
            [1, "1002", "银行存款", "2035-01", "期初余额", "", 12000.00, 5000.00, "借", 7000.00],
            [2, "3001", "实收资本", "2035-01", "期初余额", "", None, 7000.00, "贷", 7000.00],
        ]
        st, r = upload("/api/data/import/detail-ledger?book_id=t11x", "net.xlsx",
                       xlsx(rows, DETAIL_HEADERS))
        assert st == 200, r
        assert r["opening_rows"] == 2 and r["opening_balanced"] is True
        assert any("净额" in w for w in r["warnings"])
        lst = get("/api/accounts/openings/list?year=2035&book_id=t11x")[1]
        by = {x["code"]: x for x in lst["rows"]}
        assert by["1002"]["debit"] == 7000 and by["1002"]["credit"] == 0
        delete("/api/booksets/t11x")


class TestImportFormatCompat:
    def _book(self, bid):
        delete(f"/api/booksets/{bid}")
        assert post("/api/booksets", {"id": bid, "name": "格式兼容", "opening_year": "2035"})[0] == 200
        return f"book_id={bid}"

    def test_续行日期留空自动补全(self):
        q = self._book("t11d")
        rows = [
            row("记-101", "2035-01-05", "提现", "1001", "库存现金", 500, 0),
            row("记-101", "", "提现", "1002", "银行存款", 0, 500),
        ]
        st, r = upload(f"/api/data/import/vouchers?{q}", "d.xlsx", xlsx(rows))
        assert st == 200 and r["created"] == 1 and r["errors"] == []
        v = get(f"/api/vouchers?period=2035-01&{q}")[1]["rows"][0]
        assert v["source_no"] == "记-101" and v["entry_count"] == 2
        delete("/api/booksets/t11d")

    def test_零金额行跳过不致整张凭证丢失(self):
        q = self._book("t11z")
        rows = [
            row("记-102", "2035-01-31", "结转本期损益", "5001", "主营业务收入", 0, 0),
            row("记-102", "2035-01-31", "结转本期损益", "5001", "主营业务收入", 100, 0),
            row("记-102", "2035-01-31", "结转本期损益", "3103", "本年利润", 0, 100),
        ]
        st, r = upload(f"/api/data/import/vouchers?{q}", "z.xlsx", xlsx(rows))
        assert st == 200 and r["created"] == 1 and r["zero_rows"] == 1
        v = get(f"/api/vouchers?period=2035-01&{q}")[1]["rows"][0]
        assert v["entry_count"] == 2 and v["total_debit"] == 100
        assert r["carryover_marked"] == 1  # 结转损益被识别
        recs = get(f"/api/carryover/records?year=2035&{q}")[1]
        assert [x["kind"] for x in recs] == ["profit"]
        delete("/api/booksets/t11z")

    def test_手工业务摘要不被误判为结转(self):
        q = self._book("t11m")
        rows = [
            row("记-1", "2035-02-05", "期初建账-现金", "1001", "库存现金", 100, 0),
            row("记-1", "2035-02-05", "期初建账-现金", "1002", "银行存款", 0, 100),
            row("记-2", "2035-02-10", "支付办公室租金", "5602", "管理费用", 300, 0),
            row("记-2", "2035-02-10", "支付办公室租金", "1002", "银行存款", 0, 300),
        ]
        st, r = upload(f"/api/data/import/vouchers?{q}", "m.xlsx", xlsx(rows))
        assert st == 200 and r["created"] == 2
        assert r["carryover_marked"] == 0
        assert get(f"/api/carryover/records?year=2035&{q}")[1] == []
        delete("/api/booksets/t11m")


class TestMonthlyCarryoverStatus:
    def test_每月结转状态矩阵包含全部结转步骤(self):
        r = get(p("/api/carryover/monthly-status"))[1]
        kinds = r["kinds"]
        assert len(kinds) == 19  # 18 步 + 结转未分配利润含在内（STEPS 全量）
        names = [k["name"] for k in kinds]
        for n in ("结转销售成本", "计提工资", "计提折旧", "计提所得税",
                  "结转本期损益", "结转未分配利润"):
            assert n in names
        assert [k["order"] for k in kinds] == sorted(k["order"] for k in kinds)
        months = {m["period"]: m for m in r["months"]}
        assert "2034-01" in months and "2034-12" in months

    def test_月份结账状态列表附带结转状态(self):
        rows = {x["period"]: x for x in get(p("/api/carryover/periods"))[1]}
        jan = rows["2034-01"]["carryover"]
        assert jan["count"] == 2
        assert sorted(jan["kinds"]) == ["profit", "sales_cost"]
        assert sorted(jan["kind_names"]) == sorted(["结转本期损益", "结转销售成本"])
        assert jan["amount"] == 74000  # 62000 + 12000
        assert jan["profit_closed"] is True
        assert {x["source"] for x in jan["records"]} == {"import"}
        feb = rows["2034-02"]["carryover"]
        assert feb["count"] == 2 and feb["profit_closed"] is True
        dec = rows["2034-12"]["carryover"]
        assert dec["kinds"] == ["retain_profit"]

    def test_结转记录带凭证号可追溯(self):
        rows = {x["period"]: x for x in get(p("/api/carryover/periods"))[1]}
        for rec in rows["2034-01"]["carryover"]["records"]:
            assert rec["voucher_no"].startswith("记-203401-")
            assert rec["amount"] > 0

    def test_未执行结转的月份显示未结平(self):
        delete("/api/booksets/t11n")
        q = "book_id=t11n"
        assert post("/api/booksets", {"id": "t11n", "name": "未结转", "opening_year": "2035"})[0] == 200
        rows = [
            row("记-1", "2035-03-10", "销售商品", "1002", "银行存款", 1130, 0),
            row("记-1", "2035-03-10", "销售商品", "5001", "主营业务收入", 0, 1000),
            row("记-1", "2035-03-10", "销售商品", "222101", "应交增值税", 0, 130),
        ]
        st, r = upload(f"/api/data/import/vouchers?{q}", "n.xlsx", xlsx(rows))
        assert st == 200 and r["carryover_marked"] == 0
        periods = {x["period"]: x for x in get(f"/api/carryover/periods?{q}")[1]}
        c = periods["2035-03"]["carryover"]
        assert c["count"] == 0 and c["kinds"] == []
        assert c["profit_closed"] is False   # 本期损益尚未结平
        assert c["pl_net"] == 1000
        delete("/api/booksets/t11n")


class TestReverseUpdatesStatus:
    def test_反结转后月份结转状态同步更新(self):
        rows = {x["period"]: x for x in get(p("/api/carryover/periods"))[1]}
        dec = rows["2034-12"]["carryover"]
        rec = next(x for x in dec["records"] if x["kind"] == "retain_profit")
        st, r = post(p(f"/api/carryover/reverse/{rec['id']}"))
        assert st == 200, r
        rows = {x["period"]: x for x in get(p("/api/carryover/periods"))[1]}
        dec = rows["2034-12"]["carryover"]
        assert dec["count"] == 0 and dec["kinds"] == []      # 已反结转不计入
        assert dec["records"][0]["status"] == "reversed"      # 记录仍可追溯
        # 反结转后损益/利润分配回退：未分配利润仍等于本年损益
        bs = get(p("/api/reports/balance-sheet?period=2034-12"))[1]
        m = {x["name"]: x for x in bs["rows"]}
        assert bs["balanced"] is True and m["未分配利润"]["ending"] == 198000
