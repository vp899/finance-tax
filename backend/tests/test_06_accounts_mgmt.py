"""科目表管理 / 批量操作 / 导入导出 / 辅助核算 / 财务人员 / 小数位 / 数据清零 测试

会计年度 2037/2038（避免与其他套件冲突），本套件最后执行数据清零类用例。
"""
import io

import pytest
from openpyxl import Workbook, load_workbook

from conftest import acc_id, api, delete, get, make_voucher, post, put, upload


@pytest.fixture(scope="module", autouse=True)
def _fresh(server):
    """依赖 session 级服务夹具（重置数据库并启动后端）"""
    yield


def _xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


class TestAccountFields:
    def test_新增科目全字段(self):
        st, r = post("/api/accounts", {
            "code": "1201", "name": "周转材料-测试", "direction": "D",
            "category": "asset", "unit": "件", "currency": "CNY",
            "quantity_accounting": True, "aux_project": True, "aux_inventory": True,
        })
        assert st == 200
        a = next(x for x in get("/api/accounts")[1] if x["code"] == "1201")
        assert a["unit"] == "件" and a["quantity_accounting"] is True
        assert a["aux_project"] is True and a["aux_inventory"] is True
        assert a["category"] == "asset" and a["category_name"] == "资产"

    def test_添加下级科目(self):
        st, r = post("/api/accounts", {
            "code": "120101", "name": "在库周转材料", "parent_code": "1201",
            "direction": "D", "category": "asset"})
        assert st == 200
        rows = {x["code"]: x for x in get("/api/accounts")[1]}
        assert rows["1201"]["is_leaf"] is False
        assert rows["120101"]["is_leaf"] is True
        assert rows["120101"]["level"] == rows["1201"]["level"] + 1

    def test_编辑科目(self):
        aid = acc_id("120101")
        st, _ = put(f"/api/accounts/{aid}", {
            "name": "在库周转材料-甲", "unit": "箱", "aux_customer": True,
            "quantity_accounting": False})
        assert st == 200
        a = next(x for x in get("/api/accounts")[1] if x["code"] == "120101")
        assert a["name"] == "在库周转材料-甲" and a["unit"] == "箱"
        assert a["aux_customer"] is True and a["quantity_accounting"] is False

    def test_核算类型成本(self):
        st, r = post("/api/accounts", {
            "code": "120102", "name": "在用周转材料", "parent_code": "1201",
            "direction": "D", "category": "cost"})
        assert st == 200
        a = next(x for x in get("/api/accounts")[1] if x["code"] == "120102")
        assert a["category"] == "cost" and a["category_name"] == "成本"

    def test_非法核算类型被拒(self):
        st, _ = post("/api/accounts", {"code": "1299", "name": "x", "category": "bad"})
        assert st == 400

    def test_科目编码重复被拒(self):
        st, _ = post("/api/accounts", {"code": "1201", "name": "重复"})
        assert st == 400

    def test_期初余额按年月查询(self):
        aid = acc_id("120101")
        assert put("/api/accounts/openings/save", {
            "year": "2037",
            "rows": [{"account_id": aid, "debit": 1000},
                     {"account_id": acc_id("3001"), "credit": 1000}]})[0] == 200
        rows = get("/api/accounts?as_of=2037-01")[1]
        a = next(x for x in rows if x["code"] == "120101")
        assert a["opening"]["debit"] == 1000
        b = next(x for x in rows if x["code"] == "3001")
        assert b["opening"]["credit"] == 1000


class TestDisableBatch:
    def _mk(self, code, name, parent="1201"):
        st, r = post("/api/accounts", {
            "code": code, "name": name, "parent_code": parent,
            "direction": "D", "category": "asset"})
        assert st == 200
        return r["id"]

    def test_停用科目不能记账(self):
        aid = self._mk("120103", "停用测试")
        assert put(f"/api/accounts/{aid}", {"is_disabled": True})[0] == 200
        st, r = post("/api/vouchers", {
            "date": "2037-03-05",
            "entries": [{"account_id": aid, "summary": "停用", "debit": 10},
                        {"account_id": acc_id("3001"), "summary": "停用", "credit": 10}]})
        assert st == 400 and "停用" in r["detail"]

    def test_停用科目不出现在智能补全(self):
        r = get("/api/vouchers/suggest?q=停用测试")[1]["accounts"]
        assert all(a["code"] != "120103" for a in r)

    def test_批量启用(self):
        st, r = post("/api/accounts/batch", {"action": "enable", "ids": [acc_id("120103")]})
        assert st == 200 and r["ok_count"] == 1
        a = next(x for x in get("/api/accounts")[1] if x["code"] == "120103")
        assert a["is_disabled"] is False

    def test_批量禁用(self):
        ids = [acc_id("120103"), acc_id("120102")]
        st, r = post("/api/accounts/batch", {"action": "disable", "ids": ids})
        assert st == 200 and r["ok_count"] == 2
        for code in ("120103", "120102"):
            a = next(x for x in get("/api/accounts")[1] if x["code"] == code)
            assert a["is_disabled"] is True
        assert post("/api/accounts/batch", {"action": "enable", "ids": ids})[0] == 200

    def test_批量删除部分成功(self):
        aid = self._mk("120104", "待删科目")
        # 有下级/有期初的不允许删除
        st, r = post("/api/accounts/batch", {"action": "delete", "ids": [aid, acc_id("1201")]})
        res = {x["id"]: x for x in r["results"]}
        assert res[aid]["ok"] is True
        assert res[acc_id("1201")]["ok"] is False  # 有期初 1000
        assert all(x["code"] != "120104" for x in get("/api/accounts")[1])

    def test_非法批量操作被拒(self):
        assert post("/api/accounts/batch", {"action": "unknown", "ids": [1]})[0] == 400
        assert post("/api/accounts/batch", {"action": "delete", "ids": []})[0] == 400


class TestCostCategory:
    def test_成本类记账与报表覆盖(self):
        st, v = make_voucher("2037-04-05", [
            ("4001", "领用材料", 5000, 0), ("1002", "领用材料", 0, 5000)])
        assert st == 200
        tb = get("/api/books/trial-balance?period=2037-04")[1]
        assert tb["balanced"] is True
        bs = get("/api/reports/balance-sheet?period=2037-04")[1]
        assert bs["balanced"] is True
        m = {x["name"]: x for x in bs["rows"]}
        # 成本类并入存货
        assert m["存货"]["ending"] == 5000


class TestDecimalSettings:
    def test_小数位设置保存(self):
        assert put("/api/settings", {"decimal_qty": "2", "decimal_price": "2",
                                     "decimal_rate": "6"})[0] == 200
        s = get("/api/settings")[1]
        assert s["decimal_qty"] == "2" and s["decimal_rate"] == "6"

    def test_小数位非法值被拒(self):
        assert put("/api/settings", {"decimal_qty": "9"})[0] == 400
        assert put("/api/settings", {"decimal_qty": "abc"})[0] == 400

    def test_数量两位小数校验(self):
        st, r = post("/api/vouchers", {
            "date": "2037-05-05",
            "entries": [{"account_id": acc_id("120101"), "summary": "数量", "debit": 100,
                         "quantity": 1.234},
                        {"account_id": acc_id("3001"), "summary": "数量", "credit": 100}]})
        assert st == 400 and "数量最多 2 位小数" in r["detail"]

    def test_汇率六位小数校验(self):
        st, r = post("/api/vouchers", {
            "date": "2037-05-05",
            "entries": [{"account_id": acc_id("120101"), "summary": "汇率", "debit": 100,
                         "exchange_rate": 7.1234567},
                        {"account_id": acc_id("3001"), "summary": "汇率", "credit": 100}]})
        assert st == 400 and "汇率最多 6 位小数" in r["detail"]

    def test_合规小数位通过(self):
        st, r = post("/api/vouchers", {
            "date": "2037-05-06",
            "entries": [{"account_id": acc_id("120101"), "summary": "数量", "debit": 100,
                         "quantity": 1.23, "exchange_rate": 7.123456},
                        {"account_id": acc_id("3001"), "summary": "数量", "credit": 100}]})
        assert st == 200
        e = get(f"/api/vouchers/{r['id']}")[1]["entries"][0]
        assert e["quantity"] == 1.23


class TestStaffAuxSettings:
    def test_财务人员设置(self):
        assert put("/api/settings", {
            "staff_bookkeeper": "张三", "staff_reviewer": "李四",
            "staff_cashier": "王五", "staff_supervisor": "赵六"})[0] == 200
        s = get("/api/settings")[1]
        assert s["staff_bookkeeper"] == "张三"
        assert s["staff_reviewer"] == "李四"
        assert s["staff_cashier"] == "王五"
        assert s["staff_supervisor"] == "赵六"

    def test_辅助核算开关(self):
        assert put("/api/settings", {"aux_switch_project": "0",
                                     "aux_switch_customer": "1"})[0] == 200
        s = get("/api/settings")[1]
        assert s["aux_switch_project"] == "0"
        assert put("/api/settings", {"aux_switch_project": "x"})[0] == 400
        assert put("/api/settings", {"aux_switch_project": "1"})[0] == 200

    def test_默认值存在(self):
        s = get("/api/settings")[1]
        for k in ("decimal_qty", "decimal_price", "decimal_rate",
                  "aux_switch_project", "aux_switch_inventory"):
            assert k in s


class TestChartImportExport:
    HEADERS = ["科目编码", "科目名称", "方向", "默认币种", "数量核算", "计量单位",
               "项目", "客户", "供应商", "部门", "员工", "存货"]

    def test_导出科目表(self):
        st, data = api("GET", "/api/accounts/export.xlsx", raw=True)
        assert st == 200 and data[:2] == b"PK"
        ws = load_workbook(io.BytesIO(data)).active
        headers = [c.value for c in ws[1]]
        assert headers == self.HEADERS
        codes = {str(r[0]) for r in ws.iter_rows(min_row=2, values_only=True)}
        assert "1002" in codes and "1201" in codes

    def test_导入科目表新增与更新(self):
        rows = [
            ["101201", "工行存款-测试", "借", "CNY", "是", "个", "是", "否", "否", "否", "否", "否"],
            ["101202", "建行存款-测试", "借", "CNY", "否", "", "否", "否", "否", "否", "否", "否"],
            ["1002", "银行存款", "借", "CNY", "否", "", "否", "否", "否", "否", "否", "否"],
        ]
        st, r = upload("/api/accounts/import", "accounts.xlsx",
                       _xlsx(self.HEADERS, rows))
        assert st == 200
        assert r["created"] == 2 and r["updated"] == 1
        accs = {x["code"]: x for x in get("/api/accounts")[1]}
        assert accs["101201"]["parent_code"] == "1012"
        assert accs["101201"]["quantity_accounting"] is True
        assert accs["101201"]["aux_project"] is True
        assert accs["1012"]["is_leaf"] is False
        assert accs["1002"]["name"] == "银行存款"

    def test_导入科目表缺名称报错(self):
        rows = [["101209", "", "借", "CNY", "否", "", "否", "否", "否", "否", "否", "否"]]
        st, r = upload("/api/accounts/import", "accounts.xlsx",
                       _xlsx(self.HEADERS, rows))
        assert st == 400  # 全部失败时报错

    def test_导入导出科目表往返(self):
        st, data = api("GET", "/api/accounts/export.xlsx", raw=True)
        ws = load_workbook(io.BytesIO(data)).active
        rows = {str(r[0]): r for r in ws.iter_rows(min_row=2, values_only=True)}
        assert rows["101201"][2] == "借"
        assert rows["101201"][4] == "是"   # 数量核算
        assert rows["101201"][5] == "个"   # 计量单位
        assert rows["101201"][6] == "是"   # 项目
        assert rows["101201"][11] == "否"  # 存货


class TestOpeningsImportExport:
    HEADERS = ["科目编码", "科目名称", "币种", "辅助核算项",
               "项目", "项目编码", "客户", "客户编码", "供应商", "供应商编码",
               "部门", "部门编码", "员工", "员工编码", "存货", "存货编码",
               "规格型号", "计量单位", "方向",
               "期初数量", "期初余额原币", "期初余额本位币",
               "本年借方累计数量", "本年借方累计原币", "本年借方累计本位币",
               "本年贷方累计数量", "本年贷方累计原币", "本年贷方累计本位币"]

    def _row(self, code, name, direction, qty=0, orig=0, base=0, aux=None, spec=""):
        aux = aux or {}
        vals = {4: "", 5: "", 6: "", 7: "", 8: "", 9: "", 10: "", 11: "", 12: "", 13: "", 14: "", 15: ""}
        for (t, nidx, cidx) in [("project", 4, 5), ("customer", 6, 7), ("supplier", 8, 9),
                               ("dept", 10, 11), ("employee", 12, 13), ("inventory", 14, 15)]:
            if t in aux:
                vals[nidx], vals[cidx] = aux[t]
        return [code, name, "CNY", list(aux.keys())[0] if aux else "",
                vals[4], vals[5], vals[6], vals[7], vals[8], vals[9],
                vals[10], vals[11], vals[12], vals[13], vals[14], vals[15],
                spec, "个", direction, qty, orig, base,
                qty, orig, base, 0, 0, 0]

    def test_导入科目期初含辅助与累计(self):
        rows = [
            self._row("120101", "周转材料-测试", "借", qty=10, orig=1000, base=1000,
                      aux={"project": ("在建项目A", "P001")}, spec="φ10"),
            self._row("3001", "实收资本", "贷", base=1000),
        ]
        st, r = upload("/api/accounts/openings/import?year=2038", "openings.xlsx",
                       _xlsx(self.HEADERS, rows))
        assert st == 200
        assert r["imported"] == 2
        assert r["total_debit"] == 1000 and r["total_credit"] == 1000
        lst = get("/api/accounts/openings/list?year=2038")[1]
        assert lst["balanced"] is True
        m = {x["code"]: x for x in lst["rows"]}
        assert m["120101"]["debit"] == 1000 and m["120101"]["quantity"] == 10
        assert m["120101"]["ytd_debit"] == 1000  # 本年借方累计本位币

    def test_导入期初不平衡被拒(self):
        before = get("/api/accounts/openings/list?year=2039")[1]
        rows = [self._row("120101", "周转材料-测试", "借", base=500)]
        st, r = upload("/api/accounts/openings/import?year=2039", "openings.xlsx",
                       _xlsx(self.HEADERS, rows))
        assert st == 400 and "不平衡" in r["detail"]
        after = get("/api/accounts/openings/list?year=2039")[1]
        # 导入失败不留下脏数据（未录入期初的年度自动按上期连续累计）
        assert after["total_debit"] == before["total_debit"]
        assert after["total_credit"] == before["total_credit"]

    def test_导入期初科目不存在报错(self):
        rows = [self._row("9999", "不存在", "借", base=100),
                self._row("120101", "周转材料-测试", "借", base=100),
                self._row("3001", "实收资本", "贷", base=100)]
        st, r = upload("/api/accounts/openings/import?year=2040", "openings.xlsx",
                       _xlsx(self.HEADERS, rows))
        assert st == 200
        assert any("9999" in e for e in r["errors"])

    def test_导出科目期初含辅助明细(self):
        st, data = api("GET", "/api/accounts/openings/export.xlsx?year=2038", raw=True)
        assert st == 200 and data[:2] == b"PK"
        ws = load_workbook(io.BytesIO(data)).active
        headers = [c.value for c in ws[1]]
        assert headers == self.HEADERS
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        by_code = {str(r[0]): r for r in rows}
        r1201 = by_code["120101"]
        assert r1201[4] == "在建项目A" and r1201[5] == "P001"   # 项目/编码
        assert r1201[16] == "φ10"                              # 规格型号
        assert r1201[18] == "借" and r1201[20] == 1000          # 方向/原币
        assert r1201[21] == 1000                                # 本位币


class TestDataClear:
    def test_清零需要确认(self):
        st, _ = post("/api/accounts/data/clear", {})
        assert st == 400

    def test_数据清零(self):
        assert make_voucher("2037-06-05", [
            ("120101", "清零前", 5, 0), ("3001", "清零前", 0, 5)])[0] == 200
        st, r = post("/api/accounts/data/clear", {"confirm": True})
        assert st == 200 and r["cleared"]["vouchers"] > 0
        assert get("/api/vouchers?size=100")[1]["total"] == 0
        assert get("/api/carryover/records")[1] == []
        # 科目与设置保留
        assert len(get("/api/accounts")[1]) > 50
        assert "decimal_qty" in get("/api/settings")[1]
