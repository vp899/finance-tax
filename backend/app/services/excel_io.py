"""Excel 导入导出"""
import io
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from sqlalchemy.orm import Session
from ..models import Account, Voucher, VoucherEntry
from . import ledger as L
from . import vouchers as V
from . import sheet

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
HEADER_FONT = Font(bold=True)
THIN = Border(*[Side(style="thin", color="B0B0B0")] * 4)
MONEY_FMT = "#,##0.00;[Red]-#,##0.00"


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN


def _autosize(ws, max_width=40):
    for col in ws.columns:
        length = 0
        letter = col[0].column_letter
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            length = max(length, min(len(v) * 2, max_width))
        ws.column_dimensions[letter].width = max(10, length + 2)


def export_vouchers(db: Session, from_period: str, to_period: str) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "凭证明细"
    headers = ["凭证号", "日期", "期间", "凭证类型", "状态", "行号", "摘要",
               "科目编码", "科目名称", "借方金额", "贷方金额", "数量", "单位", "现金流量项目",
               "规格型号", "单价", "外币金额", "币种", "汇率",
               "项目编码", "项目", "客户编码", "客户", "供应商编码", "供应商",
               "部门编码", "部门", "员工编码", "员工", "存货编码", "存货",
               "附单据数", "制单人", "审核人"]
    ws.append(headers)
    _style_header(ws)
    acc_map = {a.id: a for a in db.query(Account).all()}
    vouchers = db.query(Voucher).filter(
        Voucher.period >= from_period, Voucher.period <= to_period
    ).order_by(Voucher.date, Voucher.voucher_no).all()
    for v in vouchers:
        for e in v.entries:
            a = acc_map[e.account_id]
            aux = _load_aux(e.aux_json)
            ws.append([v.voucher_no, v.date, v.period, v.vtype, v.status, e.line_no,
                       e.summary, a.code, a.name, e.debit, e.credit,
                       e.quantity, e.unit, e.cashflow_code or "",
                       e.spec or "", e.price or 0, e.orig_amount or 0,
                       e.currency or "CNY", e.exchange_rate or 1] +
                      _aux_columns(aux) +
                      [v.attachment_count or 0, v.maker or "", v.reviewer or ""])
    for row in ws.iter_rows(min_row=2, min_col=10, max_col=11):
        for c in row:
            c.number_format = MONEY_FMT
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------- 凭证导入（模板） ----------------

VOUCHER_HEADERS = [
    "凭证类别", "凭证号", "凭证日期", "附单据数", "摘要", "科目编码", "科目名称",
    "借方金额", "贷方金额",
    "项目编码", "项目", "客户编码", "客户", "供应商编码", "供应商",
    "部门编码", "部门", "员工编码", "员工", "存货编码", "存货",
    "规格型号", "数量", "计量单位", "单价", "外币金额", "币种", "汇率", "制单人", "审核人",
]

VOUCHER_ALIASES = {
    "凭证类别": ["凭证类型", "类别", "类型"],
    "凭证号": ["凭证编号", "凭证字"],
    "凭证日期": ["日期", "记账日期", "日期(YYYY-MM-DD)", "日期(yyyy-mm-dd)"],
    "附单据数": ["附件数", "附件张数", "附单据"],
    "摘要": ["说明"],
    "科目编码": ["科目代码", "编码"],
    "科目名称": ["科目", "账户名称"],
    "借方金额": ["借方", "借方发生额"],
    "贷方金额": ["贷方", "贷方发生额"],
    "项目编码": ["项目代码"],
    "项目": ["项目名称"],
    "客户编码": ["客户代码"],
    "客户": ["客户名称"],
    "供应商编码": ["供应商代码"],
    "供应商": ["供应商名称"],
    "部门编码": ["部门代码"],
    "部门": ["部门名称"],
    "员工编码": ["员工代码"],
    "员工": ["员工名称"],
    "存货编码": ["存货代码"],
    "存货": ["存货名称"],
    "规格型号": ["规格"],
    "数量": [],
    "计量单位": ["单位"],
    "单价": [],
    "外币金额": ["原币金额"],
    "币种": ["货币", "默认币种"],
    "汇率": ["折算率"],
    "制单人": ["录入人", "制单"],
    "审核人": ["审核"],
}

# 辅助核算维度：(编码列, 名称列, JSON key)
AUX_DIMS = [
    ("项目编码", "项目", "project"),
    ("客户编码", "客户", "customer"),
    ("供应商编码", "供应商", "supplier"),
    ("部门编码", "部门", "dept"),
    ("员工编码", "员工", "employee"),
    ("存货编码", "存货", "inventory"),
]


def _load_aux(aux_json: str) -> dict:
    if not aux_json:
        return {}
    try:
        d = json.loads(aux_json)
        return d if isinstance(d, dict) else {}
    except Exception:
        return {}


def _aux_columns(aux: dict) -> list:
    out = []
    for _c, _n, key in AUX_DIMS:
        item = aux.get(key) or {}
        out.append(item.get("code", ""))
        out.append(item.get("name", ""))
    return out


def _aux_json_from(cells: dict) -> str:
    aux = {}
    for code_col, name_col, key in AUX_DIMS:
        c = sheet.text(cells.get(code_col))
        n = sheet.text(cells.get(name_col))
        if c or n:
            aux[key] = {"code": c, "name": n}
    return json.dumps(aux, ensure_ascii=False) if aux else ""


def voucher_template() -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "凭证导入模板"
    ws.append(VOUCHER_HEADERS)
    _style_header(ws)
    ws.append(["记", "记-001", "2026-01-05", 1, "提现备用金", "1001", "库存现金",
               5000, 0, "", "", "", "", "", "", "", "", "", "", "", "",
               "", "", "", "", "", "CNY", 1, "张三", "李四"])
    ws.append(["记", "记-001", "2026-01-05", 1, "提现备用金", "1002", "银行存款",
               0, 5000, "", "", "", "", "", "", "", "", "", "", "", "",
               "", "", "", "", "", "CNY", 1, "张三", "李四"])
    ws.append(["记", "记-002", "2026-01-31", 2, "采购办公用品", "5602", "管理费用",
               300, 0, "P001", "装修项目", "", "", "", "", "D01", "办公室", "", "", "", "",
               "A4 纸", 10, "箱", 30, "", "CNY", 1, "张三", "李四"])
    ws.append(["记", "记-002", "2026-01-31", 2, "采购办公用品", "1001", "库存现金",
               0, 300, "", "", "", "", "", "", "", "", "", "", "", "",
               "", "", "", "", "", "CNY", 1, "张三", "李四"])
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def import_vouchers(db: Session, data: bytes):
    """导入凭证：按（凭证日期+凭证号）合并为一张凭证；无凭证号时按（日期+凭证类型）合并

    兼容三类表头：新模板（凭证类别 凭证号 凭证日期…）、旧模板（日期 凭证类型…）、
    系统导出的凭证明细；未识别表头时按模板默认列序解析。
    """
    rows = sheet.read_rows(data)
    mapping, start = sheet.map_headers(rows, VOUCHER_ALIASES, required=("科目编码",), min_hits=3)
    if mapping:
        has_no_col = "凭证号" in mapping
    else:
        mapping = {name: i for i, name in enumerate(VOUCHER_HEADERS)}
        start, has_no_col = 1, True
    errors, created, merged, skipped = [], 0, 0, 0
    red_rows = 0
    groups = {}
    last_no, last_date = "", ""
    for i in range(start, len(rows)):
        row = rows[i]
        if not row or all(sheet.text(c) == "" for c in row):
            continue
        cells = sheet.row_cells(mapping, row)
        if sheet.text(cells.get("科目编码")).lower() == "科目编码":
            continue  # 重复表头
        date = sheet.norm_date(cells.get("凭证日期"))
        code = sheet.text(cells.get("科目编码"))
        if not date and not code and not sheet.text(cells.get("摘要")):
            continue
        if not date or not code:
            errors.append(f"第{i + 1}行：日期或科目编码为空")
            continue
        if not L.valid_period(date[:7]):
            errors.append(f"第{i + 1}行：日期无效 {date}")
            continue
        no = sheet.text(cells.get("凭证号"))
        if has_no_col and not no and last_no and date == last_date:
            no = last_no  # 凭证号只填在首行时向下补全
        if has_no_col and not no:
            errors.append(f"第{i + 1}行：凭证号为空")
            continue
        last_no, last_date = no, date
        try:
            debit = sheet.num(cells.get("借方金额"), allow_negative=True)
            credit = sheet.num(cells.get("贷方金额"), allow_negative=True)
            qty = sheet.num(cells.get("数量"))
            price = sheet.num(cells.get("单价"))
            orig = sheet.num(cells.get("外币金额"))
            rate = sheet.num(cells.get("汇率"), default=1.0)
            attach = int(sheet.num(cells.get("附单据数"), 0))
        except ValueError as e:
            errors.append(f"第{i + 1}行：{e}")
            continue
        if debit < 0 or credit < 0:
            red_rows += 1  # 红字（负数）金额原样入库
        vtype = sheet.text(cells.get("凭证类别")) or "记"
        key = (date, no) if no else (date, vtype)
        g = groups.setdefault(key, {
            "vtype": vtype, "no": no, "attach": attach,
            "maker": sheet.text(cells.get("制单人")),
            "reviewer": sheet.text(cells.get("审核人")),
            "lines": [],
        })
        g["lines"].append({
            "row": i + 1,
            "account_code": code, "summary": sheet.text(cells.get("摘要")),
            "debit": debit, "credit": credit, "quantity": qty, "price": price,
            "orig_amount": orig, "exchange_rate": rate,
            "currency": sheet.text(cells.get("币种")) or "CNY",
            "unit": sheet.text(cells.get("计量单位")),
            "spec": sheet.text(cells.get("规格型号")),
            "aux_json": _aux_json_from(cells),
        })
    for (date, no_or_type), g in groups.items():
        entries, ok = [], True
        for ln in g["lines"]:
            acc = L.account_by_code(db, ln["account_code"])
            if not acc:
                errors.append(f"第{ln['row']}行：科目 {ln['account_code']} 不存在")
                ok = False
                continue
            entries.append({
                "account_id": acc.id, "summary": ln["summary"],
                "debit": ln["debit"], "credit": ln["credit"],
                "quantity": ln["quantity"], "unit": ln["unit"],
                "price": ln["price"], "orig_amount": ln["orig_amount"],
                "currency": ln["currency"], "exchange_rate": ln["exchange_rate"],
                "spec": ln["spec"], "aux_json": ln["aux_json"],
            })
        if not ok:
            continue
        no = g["no"]
        if no:
            existing = db.query(Voucher).filter(
                Voucher.date == date,
                ((Voucher.source_no == no) | (Voucher.voucher_no == no)),
            ).first()
            if existing:
                have = [_line_key(e) for e in existing.entries]
                need = []
                pool = list(have)
                for e in entries:
                    k = _line_key(e)
                    if k in pool:
                        pool.remove(k)
                    else:
                        need.append(e)
                if not need:
                    skipped += 1
                    continue
                max_line = max([e.line_no for e in existing.entries] or [0])
                # 先按字典试算合并后是否平衡，平衡才落库（避免删未持久化分录报错）
                total_d = L.r2(sum(e.debit for e in existing.entries) +
                               sum(L.r2(e.get("debit", 0) or 0) for e in need))
                total_c = L.r2(sum(e.credit for e in existing.entries) +
                               sum(L.r2(e.get("credit", 0) or 0) for e in need))
                if abs(total_d - total_c) >= 0.005:
                    # 合并后借贷不平：拒绝追加，保持原凭证不变
                    errors.append(f"凭证 {date} {no}：合并后借贷不平衡"
                                  f"（借 {total_d:.2f} ≠ 贷 {total_c:.2f}），已拒绝")
                    continue
                for j, e in enumerate(need, 1):
                    existing.entries.append(VoucherEntry(
                        line_no=max_line + j, account_id=e["account_id"],
                        summary=(e.get("summary") or "")[:200],
                        debit=L.r2(e.get("debit", 0) or 0),
                        credit=L.r2(e.get("credit", 0) or 0),
                        quantity=round(float(e.get("quantity", 0) or 0), 2),
                        unit=e.get("unit") or "", currency=e.get("currency") or "CNY",
                        exchange_rate=round(float(e.get("exchange_rate", 1) or 1), 6),
                        spec=(e.get("spec") or "")[:100],
                        price=L.r2(e.get("price", 0) or 0),
                        orig_amount=L.r2(e.get("orig_amount", 0) or 0),
                        aux_json=(e.get("aux_json") or "")[:2000],
                    ))
                if existing.status == "draft":
                    existing.status = "posted"
                    existing.posted_at = V.now_str()
                merged += 1
                continue
        try:
            V.create_voucher(
                db, date=date, vtype=g["vtype"], entries=entries, source="import",
                source_no=no or "", attachment_count=g["attach"],
                maker=g["maker"], reviewer=g["reviewer"],
                allow_nonleaf=True,  # 历史数据导入兼容：允许直接记到非末级科目
            )
            created += 1
        except ValueError as e:
            errors.append(f"凭证 {date} {no or g['vtype']}：{e}")
    if errors and created == 0 and merged == 0:
        raise ValueError("导入失败：" + "；".join(errors[:10]))
    return {"created": created, "merged": merged, "skipped": skipped,
            "red_rows": red_rows, "errors": errors}


def _line_key(entry) -> tuple:
    """分录指纹（用于重复导入识别/合并）"""
    if isinstance(entry, VoucherEntry):
        return (entry.account_id, (entry.summary or "").strip()[:200],
                L.r2(entry.debit), L.r2(entry.credit))
    return (entry["account_id"], (entry.get("summary") or "").strip()[:200],
            L.r2(entry.get("debit", 0) or 0), L.r2(entry.get("credit", 0) or 0))


def export_table(title: str, headers: list, rows: list, money_cols: list) -> io.BytesIO:
    """通用表导出；rows 为与 headers 对应的值列表"""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:28]
    ws.append(headers)
    _style_header(ws)
    for r in rows:
        ws.append(r)
    for idx in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=idx + 1, max_col=idx + 1):
            for c in row:
                c.number_format = MONEY_FMT
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------- 科目表 / 科目期初 导入导出 ----------------

ACCOUNT_HEADERS = ["科目编码", "科目名称", "方向", "默认币种", "数量核算", "计量单位",
                   "项目", "客户", "供应商", "部门", "员工", "存货"]

OPENING_HEADERS = ["科目编码", "科目名称", "币种", "辅助核算项",
                   "项目", "项目编码", "客户", "客户编码", "供应商", "供应商编码",
                   "部门", "部门编码", "员工", "员工编码", "存货", "存货编码",
                   "规格型号", "计量单位", "方向",
                   "期初数量", "期初余额原币", "期初余额本位币",
                   "本年借方累计数量", "本年借方累计原币", "本年借方累计本位币",
                   "本年贷方累计数量", "本年贷方累计原币", "本年贷方累计本位币"]

# (aux_type, 名称列0基, 编码列0基, 中文名, account 字段)
AUX_DIMENSIONS = [
    ("project", 4, 5, "项目", "aux_project"),
    ("customer", 6, 7, "客户", "aux_customer"),
    ("supplier", 8, 9, "供应商", "aux_supplier"),
    ("dept", 10, 11, "部门", "aux_dept"),
    ("employee", 12, 13, "员工", "aux_employee"),
    ("inventory", 14, 15, "存货", "aux_inventory"),
]


def _yn(v) -> str:
    return "是" if v else "否"


def _bool_cell(v) -> int:
    return 1 if str(v or "").strip() in ("1", "是", "Y", "y", "true", "TRUE", "✓") else 0


def _dir_cell(v) -> str:
    return "借" if str(v or "").strip() in ("借", "D", "d", "借方") else "贷"


def _dir_code(v) -> str:
    return "D" if str(v or "").strip() in ("借", "D", "d", "借方") else "C"


def _num(v, default=0.0) -> float:
    if v in (None, ""):
        return default
    try:
        return float(v)
    except (TypeError, ValueError):
        raise ValueError(f"数字格式无效：{v}")


def export_accounts(db: Session) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "科目表"
    ws.append(ACCOUNT_HEADERS)
    _style_header(ws)
    for a in db.query(Account).order_by(Account.code).all():
        ws.append([
            a.code, a.name, _dir_cell(a.direction), a.currency or "CNY",
            _yn(a.quantity_accounting), a.unit or "",
            _yn(a.aux_project), _yn(a.aux_customer), _yn(a.aux_supplier),
            _yn(a.aux_dept), _yn(a.aux_employee), _yn(a.aux_inventory),
        ])
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# 科目表模板中的辅助核算列（0 基）：项目 客户 供应商 部门 员工 存货
CHART_AUX_COLUMNS = [
    ("aux_project", 6), ("aux_customer", 7), ("aux_supplier", 8),
    ("aux_dept", 9), ("aux_employee", 10), ("aux_inventory", 11),
]


def import_accounts(db: Session, data: bytes) -> dict:
    """导入科目表：按科目编码新增/更新；上级科目按编码前缀自动识别"""
    from ..models import Account
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    errors, created, updated = [], 0, 0
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or all(v in (None, "") for v in row):
            continue
        cells = (list(row) + [None] * 12)[:12]
        code = str(cells[0] or "").strip()
        name = str(cells[1] or "").strip()
        if not code:
            errors.append(f"第{i}行：科目编码为空")
            continue
        if not name:
            errors.append(f"第{i}行：科目名称为空（{code}）")
            continue
        direction = _dir_code(cells[2])
        currency = str(cells[3] or "CNY").strip() or "CNY"
        qty_acc = _bool_cell(cells[4])
        unit = str(cells[5] or "").strip()
        aux = [_bool_cell(cells[idx]) for _f, idx in CHART_AUX_COLUMNS]
        a = L.account_by_code(db, code)
        if a:
            a.name = name
            a.direction = direction
            a.currency = currency
            a.quantity_accounting = qty_acc
            a.unit = unit
            for (field, _idx), val in zip(CHART_AUX_COLUMNS, aux):
                setattr(a, field, val)
            updated += 1
        else:
            parent_code = None
            for plen in range(len(code) - 1, 0, -1):
                cand = code[:plen]
                p = L.account_by_code(db, cand)
                if p:
                    parent_code = cand
                    break
            parent = L.account_by_code(db, parent_code) if parent_code else None
            a = Account(
                code=code, name=name, parent_code=parent_code,
                direction=direction, category=L.infer_category(code, direction),
                is_leaf=1, level=(parent.level + 1) if parent else 1,
                currency=currency, unit=unit, quantity_accounting=qty_acc,
            )
            for (field, _idx), val in zip(CHART_AUX_COLUMNS, aux):
                setattr(a, field, val)
            db.add(a)
            if parent:
                parent.is_leaf = 0
            created += 1
    db.flush()
    # 维护末级标记/层级
    all_acc = db.query(Account).all()
    for a in all_acc:
        parent = next((x for x in all_acc if x.code == a.parent_code), None)
        a.level = (parent.level + 1) if parent else 1
        a.is_leaf = 0 if any(x.parent_code == a.code for x in all_acc) else 1
    if errors and created == 0 and updated == 0:
        raise ValueError("导入失败：" + "；".join(errors[:10]))
    return {"created": created, "updated": updated, "errors": errors}


def export_openings(db: Session, year: str) -> io.BytesIO:
    from ..models import OpeningBalance, OpeningBalanceItem
    wb = Workbook()
    ws = wb.active
    ws.title = "科目期初"
    ws.append(OPENING_HEADERS)
    _style_header(ws)
    items_by_acc = {}
    for it in db.query(OpeningBalanceItem).filter(OpeningBalanceItem.year == year).all():
        items_by_acc.setdefault(it.account_id, []).append(it)
    rows = 0
    for o in db.query(OpeningBalance).filter(OpeningBalance.year == year).order_by(OpeningBalance.id).all():
        acc = db.query(Account).get(o.account_id)
        if not acc:
            continue
        items = items_by_acc.get(o.account_id) or []
        if items:
            for it in items:
                aux_vals = [""] * 12
                for t, nidx, cidx, _n, _f in AUX_DIMENSIONS:
                    if it.aux_type == t:
                        aux_vals[nidx - 4] = it.aux_name
                        aux_vals[cidx - 4] = it.aux_code
                ws.append([
                    acc.code, acc.name, o.currency or "CNY", it.aux_type or "",
                    *aux_vals, it.spec or "", it.unit or acc.unit or "",
                    _dir_cell(it.direction),
                    it.qty, it.orig_amount, it.base_amount,
                    it.ytd_debit_qty, it.ytd_debit_orig, it.ytd_debit_base,
                    it.ytd_credit_qty, it.ytd_credit_orig, it.ytd_credit_base,
                ])
                rows += 1
        else:
            direction = "D" if (o.debit or 0) >= (o.credit or 0) else "C"
            ws.append([
                acc.code, acc.name, o.currency or "CNY", "",
                *[""] * 12, "", acc.unit or "", _dir_cell(direction),
                o.quantity or 0, o.orig_amount or 0,
                L.r2((o.debit or 0) + (o.credit or 0)),
                o.ytd_debit_qty or 0, o.ytd_debit_orig or 0, o.ytd_debit or 0,
                o.ytd_credit_qty or 0, o.ytd_credit_orig or 0, o.ytd_credit or 0,
            ])
            rows += 1
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def import_openings(db: Session, year: str, data: bytes) -> dict:
    """导入科目期初（含辅助核算明细与本年累计）；按全年合并口径校验试算平衡"""
    from ..models import Currency, OpeningBalance, OpeningBalanceItem
    year = str(year or "")[:4]
    if not year.isdigit():
        raise ValueError("年份无效，请在请求中指定 year=YYYY")
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    errors, imported = [], 0
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or all(v in (None, "") for v in row):
            continue
        c = (list(row) + [None] * len(OPENING_HEADERS))[:len(OPENING_HEADERS)]
        code = str(c[0] or "").strip()
        if not code:
            errors.append(f"第{i}行：科目编码为空")
            continue
        acc = L.account_by_code(db, code)
        if not acc:
            errors.append(f"第{i}行：科目 {code} 不存在")
            continue
        if not acc.is_leaf:
            errors.append(f"第{i}行：科目 {code} 不是末级科目，不能导入期初")
            continue
        try:
            currency = str(c[2] or acc.currency or "CNY").strip() or "CNY"
            qty = _num(c[19])
            orig = _num(c[20])
            base = _num(c[21])
            ydq, ydo, ydb = _num(c[22]), _num(c[23]), _num(c[24])
            ycq, yco, ycb = _num(c[25]), _num(c[26]), _num(c[27])
        except ValueError as e:
            errors.append(f"第{i}行：{e}")
            continue
        if not c[21] and orig:
            # 本位币为空时按币种汇率折算
            cur = db.query(Currency).filter(Currency.code == currency).first()
            base = L.r2(orig * (cur.rate if cur and cur.rate else 1))
        if c[18] in (None, "") and base == 0 and orig == 0:
            errors.append(f"第{i}行：方向/期初余额为空")
            continue
        direction = _dir_code(c[18])
        amt = L.r2(abs(base))
        o = db.query(OpeningBalance).filter(
            OpeningBalance.account_id == acc.id, OpeningBalance.year == year).first()
        if not o:
            o = OpeningBalance(account_id=acc.id, year=year)
            db.add(o)
        if direction == "D":
            o.debit = L.r2((o.debit or 0) + amt)
        else:
            o.credit = L.r2((o.credit or 0) + amt)
        o.quantity = L.r2((o.quantity or 0) + (qty if direction == "D" else -qty))
        o.currency = currency
        o.orig_amount = L.r2((o.orig_amount or 0) + (orig if direction == "D" else -orig))
        o.ytd_debit = L.r2((o.ytd_debit or 0) + ydb)
        o.ytd_credit = L.r2((o.ytd_credit or 0) + ycb)
        o.ytd_debit_qty = L.r2((o.ytd_debit_qty or 0) + ydq)
        o.ytd_credit_qty = L.r2((o.ytd_credit_qty or 0) + ycq)
        o.ytd_debit_orig = L.r2((o.ytd_debit_orig or 0) + ydo)
        o.ytd_credit_orig = L.r2((o.ytd_credit_orig or 0) + yco)
        # 辅助核算明细
        aux_type = aux_name = aux_code = ""
        for t, nidx, cidx, _n, _f in AUX_DIMENSIONS:
            if c[nidx] not in (None, "") or c[cidx] not in (None, ""):
                aux_type, aux_name, aux_code = t, str(c[nidx] or ""), str(c[cidx] or "")
                break
        if aux_type or c[3] or c[16]:
            db.add(OpeningBalanceItem(
                account_id=acc.id, year=year,
                aux_type=aux_type or str(c[3] or ""), aux_name=aux_name, aux_code=aux_code,
                spec=str(c[16] or ""), unit=str(c[17] or acc.unit or ""), currency=currency,
                direction=direction, qty=qty, orig_amount=orig, base_amount=amt,
                ytd_debit_qty=ydq, ytd_debit_orig=ydo, ytd_debit_base=ydb,
                ytd_credit_qty=ycq, ytd_credit_orig=yco, ytd_credit_base=ycb,
            ))
        imported += 1
    db.flush()
    total_d = total_c = 0.0
    for o in db.query(OpeningBalance).filter(OpeningBalance.year == year).all():
        total_d, total_c = L.r2(total_d + (o.debit or 0)), L.r2(total_c + (o.credit or 0))
    if abs(total_d - total_c) >= 0.005:
        raise ValueError(
            f"导入后期初试算不平衡：借方 {total_d:.2f} ≠ 贷方 {total_c:.2f}，请检查导入文件")
    if errors and imported == 0:
        raise ValueError("导入失败：" + "；".join(errors[:10]))
    return {"imported": imported, "errors": errors,
            "total_debit": total_d, "total_credit": total_c}
