"""Excel / CSV 表格读取、表头识别通用工具（供凭证导入、明细账导入等使用）"""
import io
import re
import csv
from datetime import datetime, date as _date

from openpyxl import load_workbook

_DATE_FULL = re.compile(r"^\d{4}-\d{1,2}-\d{1,2}$")
_DATE_MONTH = re.compile(r"^\d{4}-\d{1,2}$")


def text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (_date, datetime)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def num(v, default=0.0, allow_negative: bool = False) -> float:
    if v in (None, ""):
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("，", "")
    if not s:
        return default
    neg = s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    if neg and not allow_negative:
        raise ValueError(f"金额不能为负：{v}")
    if neg:
        s = s.strip("()").lstrip("-") or "0"
    try:
        return -float(s) if neg else float(s)
    except ValueError:
        raise ValueError(f"数字格式无效：{v}")


def norm_date(v) -> str:
    """日期规范化：datetime / YYYY-MM-DD / YYYY/MM/DD / YYYYMMDD / YYYY-MM"""
    s = text(v)
    if not s:
        return ""
    if isinstance(v, (_date, datetime)):
        return v.strftime("%Y-%m-%d")
    s = s.replace("/", "-").replace(".", "-")
    if re.match(r"^\d{8}$", s):
        s = f"{s[:4]}-{s[4:6]}-{s[6:]}"
    if _DATE_FULL.match(s):
        y, m, d = s.split("-")
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
    if _DATE_MONTH.match(s):
        y, m = s.split("-")
        return f"{int(y):04d}-{int(m):02d}"
    return s


def is_full_date(s: str) -> bool:
    return bool(_DATE_FULL.match(s or ""))


def is_month_date(s: str) -> bool:
    return bool(_DATE_MONTH.match(s or ""))


def read_rows(data: bytes) -> list:
    """读取 xlsx / csv / tsv 文件为二维列表"""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
        ws = wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    except Exception:
        pass
    text_data = None
    for enc in ("utf-8-sig", "gbk", "gb18030", "latin-1"):
        try:
            text_data = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text_data is None:
        raise ValueError("无法解析文件（支持 .xlsx / .csv / .txt）")
    lines = [ln for ln in text_data.splitlines() if ln.strip()]
    if not lines:
        raise ValueError("文件内容为空")
    delim = "\t" if lines[0].count("\t") >= lines[0].count(",") else ","
    out = []
    for ln in lines:
        if delim == "\t":
            out.append([c.strip() for c in ln.split("\t")])
        else:
            out.append(next(csv.reader([ln])))
    return out


def map_headers(rows: list, aliases: dict, required=("科目编码",), min_hits: int = 3):
    """识别表头行：aliases = {标准列名: [别名…]}，返回 (列名→下标, 数据起始行下标)

    识别不到（前 5 行没有足够多的已知列名）时返回 ({}, 0)，由调用方按默认列序处理。
    """
    lookup = {}
    for canon, names in aliases.items():
        for n in list(names) + [canon]:
            lookup[str(n).strip().lower()] = canon
    for i, row in enumerate(rows[:5]):
        mapping = {}
        for j, cell in enumerate(row or []):
            key = lookup.get(text(cell).lower())
            if key and key not in mapping:
                mapping[key] = j
        if len(mapping) >= min_hits and all(k in mapping for k in required):
            return mapping, i + 1
    return {}, 0


def row_cells(mapping: dict, row: list) -> dict:
    return {name: (row[idx] if idx is not None and idx < len(row) else None)
            for name, idx in mapping.items()}
