"""标准报表 Excel 导出（小企业会计准则表样，行名与财政部表样一致）

- 利润表：项目 | 本年累计 | 第一季度 | 第二季度 | 第三季度 | 第四季度
- 现金流量表：项目 | 行次 | 本月金额 | 本年累计金额
- 资产负债表：资产（期末/年初）| 负债和所有者权益（期末/年初）双栏
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import ledger as L
from ..models import CashflowItem

MONEY_FMT = "#,##0.00;[Red]-#,##0.00"
THIN = Border(*[Side(style="thin", color="B0B0B0")] * 4)
HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")


def _finish(wb, title, widths):
    ws = wb.active
    ws.title = title
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _put(ws, row, col, value, bold=False, money=False, center=False):
    c = ws.cell(row, col, value)
    c.border = THIN
    if bold:
        c.font = Font(bold=True)
    if money:
        c.number_format = MONEY_FMT
        c.alignment = Alignment(horizontal="right")
    if center:
        c.alignment = Alignment(horizontal="center")
    return c


# ---------------- 利润表（标准表样） ----------------

# (显示名, 类型, 参数)
# 主项 rev/cost/sur/sell/admin/fin/inv/nonop_in/nonop_out/tax
# 明细 codes（借方为正）/ codes_c（贷方为正）/ interest（利息费用）
# 计算 calc_op / calc_total / calc_net
INCOME_ROWS = [
    ("一、营业收入", ("rev",)),
    (" 减：营业成本", ("cost",)),
    (" 税金及附加", ("sur",)),
    (" 其中：消费税", ("codes", ["540305"])),
    (" 营业税", ("codes", ["540306"])),
    (" 城市维护建设税", ("codes", ["540301"])),
    (" 资源税", ("codes", ["540307"])),
    (" 土地增值税", ("codes", ["540308"])),
    (" 城镇土地使用税、房产税、车船税、印花税", ("codes", ["540309", "540304"])),
    (" 教育费附加、矿产资源补偿费、排污费", ("codes", ["540302", "540303"])),
    (" 销售费用", ("sell",)),
    (" 其中：商品维修费", ("codes", ["560104"])),
    (" 广告费和业务宣传费", ("codes", ["560102"])),
    (" 管理费用", ("admin",)),
    (" 其中：开办费", ("codes", ["560210"])),
    (" 业务招待费", ("codes", ["560207"])),
    (" 研究费用", ("codes", ["560211"])),
    (" 财务费用", ("fin",)),
    (" 其中：利息费用（收入以“-”号填列）", ("interest",)),
    (" 加：投资收益", ("inv",)),
    ("二、营业利润（亏损以\"－\"号填列）", ("calc_op",)),
    (" 加：营业外收入", ("nonop_in",)),
    (" 其中：政府补助", ("codes_c", ["530101"])),
    (" 减：营业外支出", ("nonop_out",)),
    (" 其中：坏账损失", ("codes", ["571103"])),
    (" 无法收回的长期债券投资损失", ("codes", ["571104"])),
    (" 无法收回的长期股权投资损失", ("codes", ["571105"])),
    (" 自然灾害等不可抗力因素造成的损失", ("codes", ["571106"])),
    (" 税收滞纳金", ("codes", ["571107"])),
    ("三、利润总额（亏损总额以\"－\"号填列）", ("calc_total",)),
    (" 减：所得税费用", ("tax",)),
    ("四、净利润（净亏损以\"－\"号填列）", ("calc_net",)),
]

INCOME_MAIN_NAMES = {"一、营业收入", " 减：营业成本", " 税金及附加", " 销售费用",
                     " 管理费用", " 财务费用", " 加：投资收益", " 加：营业外收入",
                     " 减：营业外支出", " 减：所得税费用"}
INCOME_CALC_NAMES = {"二、营业利润（亏损以\"－\"号填列）",
                     "三、利润总额（亏损总额以\"－\"号填列）",
                     "四、净利润（净亏损以\"－\"号填列）"}


def _income_values(db, from_p, to_p):
    """一组区间内利润表各行金额（剔除结转损益凭证）"""
    p = L.pl_components(db, from_p, to_p, exclude_kind="profit")

    def codes_amt(codes, sign="D"):
        total = 0.0
        for c in codes:
            v = L._code_net(db, c, to_period=to_p, from_period=from_p, exclude_kind="profit")
            total = L.r2(total + (v if sign == "D" else -v))
        return total

    vals = {}
    for name, spec in INCOME_ROWS:
        kind = spec[0]
        if kind == "rev":
            v = p["revenue"]
        elif kind == "cost":
            v = p["cost"]
        elif kind == "sur":
            v = p["tax_surcharge"]
        elif kind == "sell":
            v = p["selling"]
        elif kind == "admin":
            v = p["admin"]
        elif kind == "fin":
            v = p["finance"]
        elif kind == "inv":
            v = p["invest_income"]
        elif kind == "nonop_in":
            v = p["nonop_income"]
        elif kind == "nonop_out":
            v = p["nonop_expense"]
        elif kind == "tax":
            v = p["income_tax"]
        elif kind == "codes":
            v = codes_amt(spec[1], "D")
        elif kind == "codes_c":
            v = codes_amt(spec[1], "C")
        elif kind == "interest":
            # 利息费用 − 利息收入（收入以负数填列）
            v = L.r2(codes_amt(["560301"], "D") - codes_amt(["560303"], "C"))
        elif kind == "calc_op":
            v = L.r2(vals["一、营业收入"] - vals[" 减：营业成本"] - vals[" 税金及附加"]
                     - vals[" 销售费用"] - vals[" 管理费用"] - vals[" 财务费用"]
                     + vals[" 加：投资收益"])
        elif kind == "calc_total":
            v = L.r2(vals["二、营业利润（亏损以\"－\"号填列）"] + vals[" 加：营业外收入"]
                     - vals[" 减：营业外支出"])
        else:  # calc_net
            v = L.r2(vals["三、利润总额（亏损总额以\"－\"号填列）"] - vals[" 减：所得税费用"])
        vals[name] = v
    return vals


def export_income_standard(db, year: str) -> io.BytesIO:
    """利润表：本年累计 + 四个季度列"""
    wb = Workbook()
    ws = wb.active
    headers = ["项目", "本年累计", "第一季度", "第二季度", "第三季度", "第四季度"]
    for i, h in enumerate(headers, 1):
        c = _put(ws, 1, i, h, bold=True, center=True)
        c.fill = HEADER_FILL
    qnames = ["第一季度", "第二季度", "第三季度", "第四季度"]
    cols = [("本年累计", f"{year}-01", f"{year}-12")]
    for q in range(4):
        cols.append((qnames[q], f"{year}-{q * 3 + 1:02d}", f"{year}-{q * 3 + 3:02d}"))
    data = {label: _income_values(db, f, t) for label, f, t in cols}

    r = 2
    for name, _spec in INCOME_ROWS:
        bold = name in INCOME_MAIN_NAMES or name in INCOME_CALC_NAMES
        _put(ws, r, 1, name, bold=bold)
        for ci, (label, _f, _t) in enumerate(cols, 2):
            _put(ws, r, ci, data[label][name], bold=bold, money=True)
        r += 1
    return _finish(wb, "利润表", [48, 16, 16, 16, 16, 16])


# ---------------- 现金流量表（标准表样） ----------------

CASHFLOW_ROWS = [
    ("一、经营活动产生的现金流量：", None),
    (" 销售产成品、商品、提供劳务收到的现金", ("in", ["101"])),
    (" 收到其他与经营活动有关的现金", ("in", ["102", "103"])),
    (" 购买原材料、商品、接受劳务支付的现金", ("out", ["201"])),
    (" 支付的职工薪酬", ("out", ["202"])),
    (" 支付的税费", ("out", ["203"])),
    (" 支付其他与经营活动有关的现金", ("out", ["204"])),
    (" 经营活动产生的现金流量净额", ("net", "op")),
    ("二、投资活动产生的现金流量：", None),
    (" 收回短期投资、长期债券投资和长期股权投资收到的现金", ("in", ["301", "305"])),
    (" 取得投资收益收到的现金", ("in", ["302"])),
    (" 处置固定资产、无形资产和其他非流动资产收回的现金净额", ("in", ["303", "304"])),
    (" 短期投资、长期债券投资和长期股权投资支付的现金", ("out", ["402", "403"])),
    (" 购建固定资产、无形资产和其他非流动资产支付的现金", ("out", ["401", "404"])),
    (" 投资活动产生的现金流量净额", ("net", "inv")),
    ("三、筹资活动产生的现金流量：", None),
    (" 取得借款收到的现金", ("in", ["502", "503"])),
    (" 吸收投资者投资收到的现金", ("in", ["501"])),
    (" 偿还借款本金支付的现金", ("out", ["601"])),
    (" 偿还借款利息支付的现金", ("out", ["604"])),
    (" 分配利润支付的现金", ("out", ["605", "602", "603"])),
    (" 筹资活动产生的现金流量净额", ("net", "fin")),
    ("四、现金净增加额", ("net", "total")),
    ("加：期初现金余额", ("cash", "begin")),
    ("五、期末现金余额", ("cash", "end")),
]

CF_OP = [" 销售产成品、商品、提供劳务收到的现金", " 收到其他与经营活动有关的现金",
         " 购买原材料、商品、接受劳务支付的现金", " 支付的职工薪酬", " 支付的税费",
         " 支付其他与经营活动有关的现金"]
CF_INV = [" 收回短期投资、长期债券投资和长期股权投资收到的现金", " 取得投资收益收到的现金",
          " 处置固定资产、无形资产和其他非流动资产收回的现金净额",
          " 短期投资、长期债券投资和长期股权投资支付的现金",
          " 购建固定资产、无形资产和其他非流动资产支付的现金"]
CF_FIN = [" 取得借款收到的现金", " 吸收投资者投资收到的现金", " 偿还借款本金支付的现金",
          " 偿还借款利息支付的现金", " 分配利润支付的现金"]
CF_SIGN = ["in", "in", "out", "out", "out", "out"]


def _net_of(store, names, signs):
    return L.r2(sum(v if s == "in" else -v for v, s in zip((store[n] for n in names), signs)))


# 表样外现金流量项目的折入行（与屏幕上的现金流量表同规则）
CF_FOLD_ROWS = {
    ("operating", "in"): " 收到其他与经营活动有关的现金",
    ("operating", "out"): " 支付其他与经营活动有关的现金",
    ("investing", "in"): " 收回短期投资、长期债券投资和长期股权投资收到的现金",
    ("investing", "out"): " 购建固定资产、无形资产和其他非流动资产支付的现金",
    ("financing", "in"): " 取得借款收到的现金",
    ("financing", "out"): " 分配利润支付的现金",
}


def _cashflow_values(db, from_p, to_p):
    """一组区间内现金流量表各行金额（本月/本年累计各自一套）"""
    amounts = L.cashflow_amounts(db, from_p, to_p)

    def item_val(code):
        inflow, outflow = amounts.get(code, (0.0, 0.0))
        return L.r2(inflow - outflow)

    # 行 → 项目编码；表样外项目折入同类别同方向的“其他”行，保证不漏项
    row_codes = {name: list(spec[1]) for name, spec in CASHFLOW_ROWS
                 if spec and spec[0] in ("in", "out")}
    known = {c for codes in row_codes.values() for c in codes}
    meta = {c.code: (c.category, c.direction) for c in db.query(CashflowItem).all()}
    for code, (inflow, outflow) in amounts.items():
        if code in known:
            continue
        target = CF_FOLD_ROWS.get(L.cashflow_bucket(code, meta, inflow, outflow))
        if target in row_codes:
            row_codes[target].append(code)

    def line_val(name, direction):
        total = 0.0
        for c in row_codes.get(name, []):
            v = item_val(c)
            total = L.r2(total + (v if direction == "in" else -v))
        return total

    cash_ids = []
    for c in L.CASH_CODES:
        cash_ids.extend(L.account_ids_for(db, c, rollup=True))

    vals, store = {}, {}
    for name, spec in CASHFLOW_ROWS:
        if spec is None:
            vals[name] = None
            continue
        kind, arg = spec
        if kind in ("in", "out"):
            v = line_val(name, kind)
        elif kind == "net":
            if arg == "op":
                v = _net_of(store, CF_OP, CF_SIGN)
            elif arg == "inv":
                v = _net_of(store, CF_INV, ["in", "in", "in", "out", "out"])
            elif arg == "fin":
                v = _net_of(store, CF_FIN, ["in", "in", "out", "out", "out"])
            else:
                v = L.r2(store[" 经营活动产生的现金流量净额"]
                         + store[" 投资活动产生的现金流量净额"]
                         + store[" 筹资活动产生的现金流量净额"])
        else:  # cash
            if arg == "begin":
                d, c = L.opening_sums(db, cash_ids, from_p)
                v = L.r2(d - c)
            else:
                v = L.r2(store["加：期初现金余额"] + store["四、现金净增加额"])
        vals[name] = v
        store[name] = v
    return vals


def export_cashflow_standard(db, period: str) -> io.BytesIO:
    """现金流量表：本月金额 + 本年累计金额"""
    year = period[:4]
    wb = Workbook()
    ws = wb.active
    headers = ["项目", "行次", "本月金额", "本年累计金额"]
    for i, h in enumerate(headers, 1):
        c = _put(ws, 1, i, h, bold=True, center=True)
        c.fill = HEADER_FILL
    month_vals = _cashflow_values(db, period, period)
    ytd_vals = _cashflow_values(db, f"{year}-01", period)

    r = 2
    for line_no, (name, spec) in enumerate(CASHFLOW_ROWS, 1):
        is_head = spec is None
        is_net = spec is not None and spec[0] in ("net", "cash")
        _put(ws, r, 1, name, bold=is_head or is_net)
        _put(ws, r, 2, line_no, center=True)
        _put(ws, r, 3, "" if is_head else month_vals[name], bold=is_net, money=not is_head)
        _put(ws, r, 4, "" if is_head else ytd_vals[name], bold=is_net, money=not is_head)
        r += 1
    return _finish(wb, "现金流量表", [52, 8, 18, 18])


# ---------------- 资产负债表（双栏表样） ----------------

def export_balance_sheet_twocol(db, period: str) -> io.BytesIO:
    """资产负债表：资产 | 期末/年初 ‖ 负债和所有者权益 | 期末/年初"""
    data = L.balance_sheet(db, period)
    left = [r for r in data["rows"] if str(r.get("group", "")).startswith("asset")]
    right = [r for r in data["rows"] if not str(r.get("group", "")).startswith("asset")]
    wb = Workbook()
    ws = wb.active
    headers = ["资产", "期末余额", "年初余额", "负债和所有者权益", "期末余额", "年初余额"]
    for i, h in enumerate(headers, 1):
        c = _put(ws, 1, i, h, bold=True, center=True)
        c.fill = HEADER_FILL

    n = max(len(left), len(right))
    for i in range(n):
        r = i + 2
        l = left[i] if i < len(left) else None
        rt = right[i] if i < len(right) else None
        lb = bool(l and l["type"] != "line")
        rb = bool(rt and rt["type"] != "line")
        _put(ws, r, 1, l["name"] if l else "", bold=lb)
        _put(ws, r, 2, l["ending"] if l else "", bold=lb, money=bool(l))
        _put(ws, r, 3, l["beginning"] if l else "", bold=lb, money=bool(l))
        _put(ws, r, 4, rt["name"] if rt else "", bold=rb)
        _put(ws, r, 5, rt["ending"] if rt else "", bold=rb, money=bool(rt))
        _put(ws, r, 6, rt["beginning"] if rt else "", bold=rb, money=bool(rt))
    return _finish(wb, "资产负债表", [32, 16, 16, 36, 16, 16])
